# -*- coding: utf-8 -*-
"""TOP SPECIES PER KINGDOM: an Excel workbook with charts and a Word report (2026-09-09).

WHAT IT IS FOR
--------------
From the delivery identity table (one row per bin, PAK name and level) and the clean-bin
tables (reads per bin, reads per barcode) it lists, for archaea (libraries A1 + A2),
bacteria (B) and fungi (F1 + F2), the N most abundant species by their share of the
barcode's FULL-LENGTH amplicon reads, per sampling year, and writes:

    <out>/TOP_SPECIES.xlsx        one sheet per kingdom (table, bar chart, per-year line chart),
                                  a method sheet with the share definition and denominators
    <out>/TOP_SPECIES_REPORT.docx method, one section per kingdom with table and charts, limits
    <out>/charts/*.png

RULES THAT KEEP THE LIST HONEST
-------------------------------
* Only species-level rows count as species (level SPECIES or CANDIDATE 'cf.'); the same
  species named 'cf.' in one bin and plainly in another is one row, shown as 'cf.' only
  when no bin reaches the plain species level. Genus-level rows fill the list up to N
  only when species run out, marked "(genus level)" in grey italics: they are not species.
* KINGDOMS DO NOT MIX. Every name is mapped to its kingdom from the reference headers
  (SILVA SSU path, then GTDB, then the fungal ITS sets); a name whose kingdom is not the
  library's (a bacterium amplified in the archaeal library, a ciliate in the fungal one)
  is listed separately as "off-target" and never enters a kingdom's ranking. Bins the
  identity step marked "outside kingdom" enter no list.
* Share = (full-length reads + concatemer segments of the bin) / (full-length amplicon
  reads of the barcode). Fragments and short reads are not in the denominator: they
  cannot enter a bin or an identity. The share against ALL reads of the barcode is kept
  in a column for information.

RUN
---
    python3 steps/top_species_report.py --root <study root> --bins <clean-bin root> [--out DIR] [--n 20]
        [--refdb ~/refdb_hizli]

--root must hold TESLIM_2026-08-12/3_VERI/KIMLIK_EN_YUKSEK.tsv (the delivery table) and
"fastq files/SELECTION_TABLE.tsv" or "fastq files/SECIM_TABLOSU.tsv" (the bin selection);
--bins holds <barcode>/BIN_TABLE.tsv or OBEK_TABLOSU.tsv (clean_bins / KUTU_TEMIZ output).
"""
from __future__ import print_function
import argparse
import collections
import io
import os
import subprocess
import sys

KINGDOMS = [(u'Archaea', ('A1', 'A2')), (u'Bacteria', ('B',)), (u'Fungi', ('F1', 'F2'))]
YEAR = {u'1': u'2021', u'2': u'2022', u'3': u'2024', u'4': u'2025'}
SPECIES_LEVELS = (u'TÜR', u'TÜR ADAYI (cf.)', u'SPECIES', u'CANDIDATE (cf.)')
GENUS_LEVELS = (u'CİNS', u'GENUS')
KINGDOM_OF = {u'Archaea': u'Archaea', u'Bacteria': u'Bacteria', u'Eukaryota': u'Fungi', u'Fungi': u'Fungi'}


def tsv(path):
    if not os.path.exists(path):
        return []
    with io.open(path, encoding='utf-8', errors='replace') as fh:
        return [s.rstrip(u'\n').split(u'\t') for s in fh if s.strip()]


def col(head, *names):
    for n in names:
        if n in head:
            return head.index(n)
    raise SystemExit('column not found: %s (have: %s)' % (' / '.join(names), ' | '.join(head)))


def kingdom_map(names, refdb, cache):
    """name -> kingdom from the reference headers; 'other eukaryote' for Eukaryota without Fungi; '?' unknown."""
    out = {}
    for r in tsv(cache)[1:]:
        out[r[0]] = r[1]
    silva = os.path.join(refdb, 'SILVA_138.2_SSURef_NR99.fasta')
    gtdb = os.path.join(refdb, 'GTDB_ssu_all_r220.fna')
    fungal = [os.path.join(refdb, 'UNITE_ITS.fasta'), os.path.join(refdb, 'fungi.ITS.fna')]
    new = False
    for name in sorted(set(names)):
        if name in out or not name:
            continue
        p = name.replace(u'cf. ', u'').split()
        key = u' '.join(p[:2]) if p and p[0] == u'Candidatus' else (p[0] if p else u'')
        k = u'?'
        if key and os.path.exists(silva):
            for pattern in (u';%s;' % key, u';%s' % key):
                try:
                    line = subprocess.check_output(['grep', '-m1', '-F', pattern, silva]).decode('utf-8', 'replace')
                    path = line.split(None, 1)[1]
                    k = KINGDOM_OF.get(path.split(u';')[0].strip(), u'?')
                    if k == u'Fungi' and u';Fungi' not in path:
                        k = u'other eukaryote'
                    break
                except subprocess.CalledProcessError:
                    continue
        if k == u'?' and key and os.path.exists(gtdb):
            try:
                line = subprocess.check_output(['grep', '-m1', '-F', u'g__%s;' % key, gtdb]).decode('utf-8', 'replace')
                k = KINGDOM_OF.get(line.split(u'd__', 1)[1].split(u';')[0].strip(), u'?') if u'd__' in line else u'?'
            except subprocess.CalledProcessError:
                pass
        if k == u'?' and key:
            for f in fungal:
                if os.path.exists(f) and subprocess.call(['grep', '-q', '-F', key, f]) == 0:
                    k = u'Fungi'
                    break
        out[name] = k
        new = True
    if new:
        io.open(cache, 'w', encoding='utf-8', newline='\n').write(
            u'name\tkingdom\n' + u'\n'.join(u'%s\t%s' % (a, b) for a, b in sorted(out.items())) + u'\n')
    return out


def bin_shares(root, bins_root):
    """bin -> (reads, share of full-length reads %, share of all reads %); denominators per barcode."""
    sel = tsv(os.path.join(root, 'fastq files', 'SELECTION_TABLE.tsv')) or \
        tsv(os.path.join(root, 'fastq files', 'SECIM_TABLOSU.tsv'))
    if not sel:
        raise SystemExit('no selection table under "fastq files"')
    h = sel[0]
    ib, ibin, ir, ish, ic = (col(h, u'barcode', u'barkod'), col(h, u'bin', u'obek'), col(h, u'reads', u'okuma'),
                             col(h, u'share %', u'pay %'), col(h, u'chosen', u'secildi'))
    denom, out = {}, {}
    for r in sel[1:]:
        if r[ic] not in (u'YES', u'EVET'):
            continue
        bc = r[ib]
        if bc not in denom:
            t = tsv(os.path.join(bins_root, bc, 'BIN_TABLE.tsv')) or tsv(os.path.join(bins_root, bc, 'OBEK_TABLOSU.tsv'))
            full = seg = 0
            if t:
                th = t[0]
                i_full, i_seg = col(th, u'reads', u'okuma (tam)'), col(th, u'concatemer segments', u'konkatemer segmenti')
                for x in t[1:]:
                    if x[0].startswith((u'BIN', u'OBEK')):
                        full += int(x[i_full])
                        seg += int(x[i_seg])
                    elif x[0].startswith((u'full unassigned', u'small cluster', u'tam atanmadi', u'kucuk obek')):
                        full += int(x[1])
            denom[bc] = max(1, full + seg)
        reads = int(r[ir])
        out[u'%s_%s' % (bc, r[ibin])] = (reads, 100.0 * reads / denom[bc], float(r[ish]))
    return out, denom


def collect(root, bins_root, n, refdb):
    k = tsv(os.path.join(root, 'TESLIM_2026-08-12', '3_VERI', 'KIMLIK_EN_YUKSEK.tsv'))
    h = k[0]
    ik, iname, ilev, ipid, iloc = (col(h, u'kutu', u'bin'), col(h, u'SÜZGEÇTEN GEÇEN AD', u'name'),
                                  col(h, u'düzey', u'level'), col(h, u'özdeşlik %', u'identity %'), col(h, u'lokus', u'locus'))
    int_ = h.index(u'NCBI nt') if u'NCBI nt' in h else None
    rows = [r for r in k[1:] if r and r[0] and not r[0].startswith(u'#') and len(r) > ilev]
    share, denom = bin_shares(root, bins_root)
    kmap = kingdom_map([r[iname].strip() for r in rows if r[ilev].strip() in SPECIES_LEVELS + GENUS_LEVELS],
                       refdb, os.path.join(root, 'GUNCEL', 'KINGDOM_MAP.tsv'))
    result = collections.OrderedDict()
    for kingdom, groups in KINGDOMS:
        samples = [u'%s-%s' % (g, y) for g in groups for y in (u'1', u'2', u'3', u'4')]
        taxa = collections.defaultdict(lambda: dict(shares=collections.Counter(), all_reads=collections.Counter(),
                                                    bins=[], identity=0.0, loci=set(), reads=0, species_bins=0, nt=set()))
        offtarget = collections.defaultdict(lambda: dict(shares=collections.Counter(), bins=[], kingdom=u'?', level=u''))
        others = {}
        for r in rows:
            bin_id = r[ik]
            if bin_id.split(u'-')[0] not in groups or bin_id not in share:
                continue
            name, level = r[iname].strip(), r[ilev].strip()
            sample = bin_id.split(u'_')[0]
            reads, s_full, s_all = share[bin_id]
            if level in SPECIES_LEVELS + GENUS_LEVELS:
                kk = kmap.get(name, u'?')
                if kk != kingdom:
                    o = offtarget[name]
                    o['shares'][sample] += s_full
                    o['bins'].append(bin_id)
                    o['kingdom'], o['level'] = kk, level
                    continue
            if level in SPECIES_LEVELS:
                key = (name.replace(u' cf. ', u' '), u'species')
            elif level in GENUS_LEVELS:
                key = (name, u'genus level')
            else:
                others[(sample, name, level)] = s_full
                continue
            e = taxa[key]
            if level in (u'TÜR', u'SPECIES'):
                e['species_bins'] += 1
            e['shares'][sample] += s_full
            e['all_reads'][sample] += s_all
            e['bins'].append(bin_id)
            e['reads'] += reads
            e['loci'].add(r[iloc])
            if int_ is not None and len(r) > int_ and r[int_] not in (u'', u'-'):
                e['nt'].add(r[int_][:60])
            try:
                e['identity'] = max(e['identity'], float(r[ipid].replace(u',', u'.')))
            except ValueError:
                pass

        def mean(e):
            return sum(e['shares'].values()) / float(len(samples))
        species = sorted([(a, e) for a, e in taxa.items() if a[1] == u'species'], key=lambda x: -mean(x[1]))
        genera = sorted([(a, e) for a, e in taxa.items() if a[1] == u'genus level'], key=lambda x: -mean(x[1]))
        chosen = [((a[0], u'species' if e['species_bins'] else u'cf.'), e) for a, e in species[:n]]
        if len(chosen) < n:
            chosen += genera[:n - len(chosen)]
        result[kingdom] = dict(samples=samples, rows=chosen, n_species=len(species), n_genera=len(genera), mean=mean,
                               species_share=sum(sum(e['shares'].values()) for _a, e in species) / float(len(samples)),
                               genus_share=sum(sum(e['shares'].values()) for _a, e in genera) / float(len(samples)),
                               top_others=sorted(others.items(), key=lambda x: -x[1])[:6],
                               offtarget=sorted(offtarget.items(), key=lambda x: -sum(x[1]['shares'].values())))
    return result, denom, len(rows)


def label(a):
    name, kind = a
    if kind == u'genus level':
        return name + u' (genus level)'
    if kind == u'cf.' and u' ' in name:
        return u'%s cf. %s' % tuple(name.split(u' ', 1))
    return name


def write_excel(result, denom, path, n):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    hfont, hfill = Font(bold=True, color='FFFFFF', name='Arial', size=10), PatternFill('solid', fgColor='1F3864')
    for kingdom, s in result.items():
        ws = wb.create_sheet(u'%s top %d' % (kingdom, n))
        ws.append([u'rank', u'species', u'level', u'mean share % (full-length)', u'max share %'] +
                  [u'%s %s' % (o.split(u'-')[0], YEAR[o.split(u'-')[1]]) for o in s['samples']] +
                  [u'bins', u'bin ids', u'reads', u'max identity %', u'loci', u'mean share % (all reads)', u'NCBI nt'])
        for c in ws[1]:
            c.font, c.fill, c.alignment = hfont, hfill, Alignment(wrap_text=True, vertical='center')
        for i, (a, e) in enumerate(s['rows'], 1):
            ws.append([i, label(a), a[1], round(s['mean'](e), 3), round(max(e['shares'].values()) if e['shares'] else 0, 3)] +
                      [round(e['shares'].get(o, 0.0), 3) for o in s['samples']] +
                      [len(e['bins']), u', '.join(e['bins']), e['reads'], round(e['identity'], 2), u'/'.join(sorted(e['loci'])),
                       round(sum(e['all_reads'].values()) / float(len(s['samples'])), 3), u'; '.join(sorted(e['nt'])) or u'-'])
            for c in ws[ws.max_row]:
                c.font = Font(italic=True, color='777777', name='Arial', size=10) if a[1] == u'genus level' else Font(name='Arial', size=10)
        ws.column_dimensions['B'].width = 42
        ws.column_dimensions[get_column_letter(7 + len(s['samples']))].width = 40
        for c in range(3, 6 + len(s['samples'])):
            ws.column_dimensions[get_column_letter(c)].width = 11
        ws.freeze_panes = 'C2'
        m = len(s['rows'])
        note = ws.cell(row=m + 3, column=1, value=u'Kingdom summary (mean over samples): species-level share %.2f%%; genus-level %.2f%%; '
                       u'largest bins without a species name: %s' % (s['species_share'], s['genus_share'], u'; '.join(
                           u'%s %s (%s) %.1f%%' % (k[0], k[1][:36], k[2][:14], v) for k, v in s['top_others'])))
        note.font = Font(italic=True, name='Arial', size=9, color='555555')
        if s['offtarget']:
            o = ws.cell(row=m + 4, column=1, value=u'Off-target (another kingdom amplified in this library; not ranked): ' + u'; '.join(
                u'%s [%s, %s] mean share %.2f%% (%s)' % (nm, e['kingdom'], e['level'], sum(e['shares'].values()) / float(len(s['samples'])), u', '.join(e['bins']))
                for nm, e in s['offtarget']))
            o.font = Font(italic=True, name='Arial', size=9, color='8B0000')
        if m:
            ch = BarChart()
            ch.type, ch.style, ch.legend = 'bar', 10, None
            ch.title = u'%s: top %d, mean full-length share (%%)' % (kingdom, m)
            ch.add_data(Reference(ws, min_col=4, min_row=1, max_row=m + 1), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=2, min_row=2, max_row=m + 1))
            ch.height, ch.width = 0.6 * m + 4, 22
            ws.add_chart(ch, 'A%d' % (m + 6))
            lc = LineChart()
            lc.title = u'%s: top 10 by year (%%)' % kingdom
            for r in range(2, min(m, 10) + 2):
                lc.add_data(Reference(ws, min_col=6, max_col=5 + len(s['samples']), min_row=r, max_row=r), from_rows=True, titles_from_data=False)
                lc.series[-1].title = SeriesLabel(v=ws.cell(row=r, column=2).value[:40])
            lc.set_categories(Reference(ws, min_col=6, max_col=5 + len(s['samples']), min_row=1, max_row=1))
            lc.height, lc.width = 12, 22
            ws.add_chart(lc, 'A%d' % (m + 6 + int(0.6 * m + 4) * 2 + 2))
    ws = wb.create_sheet(u'Method')
    for row in [
        [u'Source', u'PAK identity table (KIMLIK_EN_YUKSEK.tsv) on clean bins (one organism, full-length amplicons)'],
        [u'Kingdoms', u'Archaea = A1 + A2, Bacteria = B, Fungi = F1 + F2; names mapped to kingdoms from SILVA / GTDB / UNITE headers; off-target names listed apart'],
        [u'Years', u'-1 = 2021, -2 = 2022, -3 = 2024, -4 = 2025'],
        [u'Species row', u'level SPECIES or cf.; "cf." = the two nearest references sit within 0.5 points, the name is a candidate'],
        [u'(genus level)', u'added only when fewer than N species exist; italic and marked; not a species'],
        [u'Share % (full-length)', u'bin reads (full-length + cut concatemer segments) / full-length amplicon reads of the barcode'],
        [u'Share % (all reads)', u'the same against every read of the barcode, for information'],
        [u'Thresholds', u'16S species 98.7 / genus 94.5 (Kim 2014, Yarza 2014); ITS 99.6 / 94.3, 28S 99.8 / 98.2 (Vu 2018); margin 0.5'],
        [u'Denominators', u'; '.join(u'%s %d' % (k, v) for k, v in sorted(denom.items()))],
    ]:
        ws.append(row)
    ws.column_dimensions['A'].width, ws.column_dimensions['B'].width = 24, 140
    wb.save(path)


def write_charts(result, folder):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    if not os.path.isdir(folder):
        os.makedirs(folder)
    files = {}
    for kingdom, s in result.items():
        rows = s['rows']
        if not rows:
            continue
        names = [label(a) for a, _e in rows][::-1]
        vals = [s['mean'](e) for _a, e in rows][::-1]
        colors = ['#9e9e9e' if a[1] == u'genus level' else '#1f3864' for a, _e in rows][::-1]
        fig, ax = plt.subplots(figsize=(9, 0.32 * len(rows) + 1.6))
        ax.barh(names, vals, color=colors)
        ax.set_xlabel(u'mean full-length amplicon share (%)')
        ax.set_title(u'%s: top %d (PAK, clean bins)' % (kingdom, len(rows)))
        for i, v in enumerate(vals):
            ax.text(v, i, u' %.2f' % v, va='center', fontsize=7)
        plt.tight_layout()
        p1 = os.path.join(folder, u'%s_top.png' % kingdom.lower())
        fig.savefig(p1, dpi=160)
        plt.close(fig)
        fig, ax = plt.subplots(figsize=(9, 4.2))
        ticks = [u'%s %s' % (o.split(u'-')[0], YEAR[o.split(u'-')[1]]) for o in s['samples']]
        for a, e in rows[:10]:
            ax.plot(ticks, [e['shares'].get(o, 0.0) for o in s['samples']], marker='o', linewidth=1.2, label=label(a)[:36])
        ax.set_ylabel(u'share (%)')
        ax.set_title(u'%s: top 10 by year' % kingdom)
        ax.tick_params(axis='x', labelrotation=30, labelsize=8)
        ax.legend(fontsize=7, ncol=2)
        plt.tight_layout()
        p2 = os.path.join(folder, u'%s_years.png' % kingdom.lower())
        fig.savefig(p2, dpi=160)
        plt.close(fig)
        files[kingdom] = (p1, p2)
    return files


def write_word(result, charts, path, n, n_bins, date):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    NAVY, GREY = RGBColor(0x1F, 0x38, 0x64), RGBColor(0x55, 0x55, 0x55)

    def shade(cell, colour):
        el = OxmlElement('w:shd')
        el.set(qn('w:val'), 'clear')
        el.set(qn('w:fill'), colour)
        cell._tc.get_or_add_tcPr().append(el)

    def table(doc, heads, rows, widths):
        t = doc.add_table(rows=1, cols=len(heads))
        t.style, t.alignment, t.autofit = 'Table Grid', WD_TABLE_ALIGNMENT.CENTER, False
        for j, b in enumerate(heads):
            c = t.rows[0].cells[j]
            c.text = ''
            r = c.paragraphs[0].add_run(b)
            r.font.bold, r.font.size, r.font.color.rgb = True, Pt(8), RGBColor(0xFF, 0xFF, 0xFF)
            shade(c, '1F3864')
        for row in rows:
            cells = t.add_row().cells
            for j, v in enumerate(row):
                cells[j].text = ''
                cells[j].paragraphs[0].add_run(u'%s' % v).font.size = Pt(8)
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)

    def heading(doc, text, level=1):
        h = doc.add_heading(text, level=level)
        for r in h.runs:
            r.font.color.rgb, r.font.name = NAVY, 'Arial'

    def para(doc, text, small=False):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.name, r.font.size = 'Arial', Pt(9 if small else 10.5)
        if small:
            r.font.color.rgb = GREY
        p.paragraph_format.space_after = Pt(6)
        return p

    def bullet(doc, text):
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(text)
        r.font.name, r.font.size = 'Arial', Pt(10.5)

    d = Document()
    d.styles['Normal'].font.name, d.styles['Normal'].font.size = 'Arial', Pt(10.5)
    for s in d.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.0)
    t = d.add_heading(u'Top %d species: archaea, bacteria and fungi' % n, level=0)
    for r in t.runs:
        r.font.color.rgb = NAVY
    p = d.add_paragraph()
    r = p.add_run(u'PAK identification on clean bins, four sampling years (2021, 2022, 2024, 2025)')
    r.font.size, r.font.color.rgb = Pt(12), GREY
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(d, date, small=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading(d, u'1. What was done', 1)
    para(d, u'Oxford Nanopore long-read amplicons of five libraries (A1, A2 archaeal 16S; B bacterial 16S; F1, F2 the '
            u'fungal rRNA operon 18S + ITS + 28S) were binned by organism without any database label (one organism per '
            u'bin, full-length amplicons only; fragments and end-to-end ligated amplicons kept apart) and %d bins were '
            u'identified with PAK (population-resolved amplicon keying).' % n_bins)
    bullet(d, u'PAK: the reads of a bin are split into populations on the locus window; the dominant population is polished '
              u'from its medoid read with minimap2 + samtools + medaka and asked against thirteen rRNA databases.')
    bullet(d, u'A species name is given only above the literature thresholds: 16S species 98.7 / genus 94.5 (Kim et al. 2014; '
              u'Yarza et al. 2014), ITS 99.6 / 94.3 and 28S 99.8 / 98.2 (Vu et al. 2018). When the two nearest references sit '
              u'within 0.5 points the name is a candidate ("cf."); "sp." is not a species.')
    bullet(d, u'Kingdoms do not mix: a record outside Fungi in the fungal library (ciliates, amoebae) receives no fungal name; '
              u'a bacterium amplified in the archaeal library is listed as off-target, never as an archaeon.')
    para(d, u'Share = reads of the bin (full-length + cut concatemer segments) / full-length amplicon reads of the barcode. '
            u'Fragments are not in the denominator. The mean is taken over the kingdom\'s samples (library x 4 years); a year '
            u'without the species counts as 0.', small=True)
    for k, (kingdom, s) in enumerate(result.items(), 2):
        heading(d, u'%d. %s: top %d' % (k, kingdom, n), 1)
        n_gen = sum(1 for a, _e in s['rows'] if a[1] == u'genus level')
        para(d, u'Species named at species level: %d (species or cf.). %s' % (
            s['n_species'], (u'%d genus-level rows were added because fewer than %d species exist; they are not species and are '
                             u'marked.' % (n_gen, n)) if n_gen else u''))
        heads = [u'#', u'species', u'level', u'mean %'] + [u'%s %s' % (o.split(u'-')[0], YEAR[o.split(u'-')[1]]) for o in s['samples']] + [u'bins', u'identity %']
        rows = [[i, label(a), a[1], u'%.2f' % s['mean'](e)] + [u'%.2f' % e['shares'].get(o, 0.0) for o in s['samples']] +
                [len(e['bins']), u'%.1f' % e['identity']] for i, (a, e) in enumerate(s['rows'], 1)]
        table(d, heads, rows, [0.8, 5.2, 1.4, 1.5] + [1.15] * len(s['samples']) + [1.0, 1.3])
        para(d, u'')
        para(d, u'Kingdom summary (mean over samples): species-level bins %.2f%%, genus-level %.2f%%. Largest bins without a '
                u'species name: %s.' % (s['species_share'], s['genus_share'], u'; '.join(u'%s %s (%s) %.1f%%' % (x[0], x[1], x[2], v) for x, v in s['top_others'])), small=True)
        if s['offtarget']:
            para(d, u'Off-target names (another kingdom amplified by this library\'s primers; valid identities, not in the %s list): %s.'
                    % (kingdom.lower(), u'; '.join(u'%s [%s, %s] mean share %.2f%%' % (nm, e['kingdom'], e['level'], sum(e['shares'].values()) / float(len(s['samples'])))
                                                    for nm, e in s['offtarget'])), small=True)
        if kingdom in charts:
            d.add_picture(charts[kingdom][0], width=Cm(16.5))
            d.add_picture(charts[kingdom][1], width=Cm(16.5))
    heading(d, u'%d. Limits' % (len(result) + 2), 1)
    bullet(d, u'"cf." names are candidates: two references within 0.5 points; better basecalling does not separate them.')
    bullet(d, u'Many bacterial bins stay at family or order level: high identity, no named species or genus in the databases. They are not in this list; the full table is in the Excel appendix.')
    bullet(d, u'Only chosen bins carry a share (>= 500 full-length reads or the five largest of a barcode); smaller clusters were not deleted and sit in the selection table with their reason.')
    heading(d, u'%d. Appendix' % (len(result) + 3), 1)
    para(d, u'TOP_SPECIES.xlsx: per-kingdom tables (year columns, bin ids, reads, identity, loci, share against all reads, NCBI nt '
            u'cross-check), bar and year charts, the method sheet with the denominators. KIMLIK_EN_YUKSEK.tsv: the full identity table of %d bins.' % n_bins)
    d.save(path)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--root', required=True)
    ap.add_argument('--bins', required=True, help='clean_bins output root (<barcode>/BIN_TABLE.tsv)')
    ap.add_argument('--out', default=None)
    ap.add_argument('--n', type=int, default=20)
    ap.add_argument('--refdb', default=os.path.expanduser('~/refdb_hizli'))
    ap.add_argument('--date', default=None)
    a = ap.parse_args(argv)
    out = a.out or os.path.join(a.root, 'GUNCEL', 'TESLIM')
    if not os.path.isdir(out):
        os.makedirs(out)
    import datetime
    date = a.date or datetime.date.today().strftime('%Y-%m-%d')
    result, denom, n_bins = collect(a.root, a.bins, a.n, a.refdb)
    for kingdom, s in result.items():
        print('  %-9s species %3d (genus-level %3d) -> listed %2d; top: %s' % (
            kingdom, s['n_species'], s['n_genera'], len(s['rows']),
            u'; '.join(u'%s %.2f%%' % (label(a)[:30], s['mean'](e)) for a, e in s['rows'][:3])))
    charts = write_charts(result, os.path.join(out, 'charts'))
    xl = os.path.join(out, 'TOP_SPECIES.xlsx')
    write_excel(result, denom, xl, a.n)
    dx = os.path.join(out, 'TOP_SPECIES_REPORT.docx')
    write_word(result, charts, dx, a.n, n_bins, date)
    print('  written: %s, %s' % (xl, dx))
    return 0


if __name__ == '__main__':
    sys.exit(main())
